from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from .processing import summary_counts
from .storage import append_log, load_state, project_dir, resolve_project_file, save_state


SCALAR_PLACEHOLDERS = {
    "{{PROJEKTNAME}}": ("project", "name"),
    "{{AUFTRAGGEBER}}": ("project", "client"),
    "{{OBJEKTADRESSE}}": ("project", "object_address"),
    "{{BESICHTIGUNGSDATUM}}": ("project", "inspection_date"),
    "{{SACHVERSTAENDIGER}}": ("project", "expert"),
    "{{SACHBEARBEITER}}": ("project", "case_worker"),
    "{{OBJEKTTYP}}": ("project", "object_type"),
}


def replace_placeholders_in_paragraph(paragraph: Any, replacements: dict[str, str]) -> bool:
    text = paragraph.text
    changed = False
    for placeholder, value in replacements.items():
        if placeholder in text:
            text = text.replace(placeholder, value)
            changed = True
    if changed:
        paragraph.text = text
    return changed


def fill_table(table: Any, master: list[dict[str, Any]]) -> None:
    headers = [
        "Proben-\nNr.",
        "Position",
        "Bohrtiefe\n[mm]",
        "Hoehe\n[cm]",
        "%Cl(Zement)\nMasse-%",
        "Karbonatisierung\n[Tiefe in mm]",
        "Bewertung",
    ]
    for col, header in enumerate(headers):
        table.cell(0, col).text = header
    for row_index, sample in enumerate(master, start=1):
        cells = table.rows[row_index].cells
        values = [
            sample.get("probe", ""),
            sample.get("position", ""),
            sample.get("drill_depth_mm", ""),
            sample.get("height_cm", ""),
            sample.get("chloride_display", ""),
            sample.get("carbonation", ""),
            sample.get("rating", ""),
        ]
        for col, value in enumerate(values):
            cells[col].text = str(value)


def add_master_table(doc: Document, master: list[dict[str, Any]], after_paragraph: Any | None = None) -> None:
    if after_paragraph is None:
        doc.add_paragraph("Master-Probenliste")
        table = doc.add_table(rows=len(master) + 1, cols=7)
    else:
        table = doc.add_table(rows=len(master) + 1, cols=7)
        after_paragraph._p.addnext(table._tbl)
    table.style = "Table Grid"
    fill_table(table, master)


def add_summary(doc: Document, state: dict[str, Any]) -> None:
    master = state.get("master_samples", [])
    counts = summary_counts(master)
    heading = doc.add_heading("Automatisch vorbereiteter Gutachtenentwurf", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    project = state.get("project", {})
    meta = doc.add_paragraph()
    meta.add_run("Projekt: ").bold = True
    meta.add_run(project.get("name", ""))
    meta.add_run("\nObjektadresse: ").bold = True
    meta.add_run(project.get("object_address", ""))
    meta.add_run("\nAuftraggeber: ").bold = True
    meta.add_run(project.get("client", ""))

    doc.add_paragraph(
        "Hinweis: Dieses Dokument ist ein pruefbarer Entwurf. Fachliche "
        "Schlussbewertung und Freigabe bleiben beim Sachverstaendigen."
    )
    doc.add_paragraph(
        f"In Summe wurden {counts['total']} Proben in der Masterliste gefuehrt. "
        f"Bei {counts['chloride_over_1']} Proben liegt der Chloridwert ueber 1,0 %. "
        f"{counts['needs_review']} Proben enthalten offene Pruefpunkte."
    )


def create_from_scratch(path: Path, state: dict[str, Any]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    add_summary(doc, state)
    add_master_table(doc, state.get("master_samples", []))
    doc.save(path)


def create_from_template(template_path: Path, output: Path, state: dict[str, Any]) -> dict[str, Any]:
    doc = Document(str(template_path))
    project = state.get("project", {})
    replacements = {
        placeholder: str(project.get(key, ""))
        for placeholder, (_, key) in SCALAR_PLACEHOLDERS.items()
    }
    replacements["{{ZUSAMMENFASSUNG_CHLORID}}"] = chloride_summary_text(state)

    inserted_table = False
    for paragraph in doc.paragraphs:
        if "{{CHLORID_TABELLE}}" in paragraph.text or "{{MASTER_PROBENLISTE}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{CHLORID_TABELLE}}", "").replace("{{MASTER_PROBENLISTE}}", "")
            add_master_table(doc, state.get("master_samples", []), after_paragraph=paragraph)
            inserted_table = True
        else:
            replace_placeholders_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_placeholders_in_paragraph(paragraph, replacements)

    if not inserted_table:
        doc.add_page_break()
        add_summary(doc, state)
        add_master_table(doc, state.get("master_samples", []))
    doc.save(output)
    return {"inserted_table_at_placeholder": inserted_table}


def chloride_summary_text(state: dict[str, Any]) -> str:
    counts = summary_counts(state.get("master_samples", []))
    return (
        f"In Summe wurden {counts['total']} Proben ausgewertet. "
        f"Bei {counts['chloride_over_1']} Proben wurde ein Chloridwert ueber 1,0 % erkannt. "
        f"Bei {counts['chloride_between_02_1']} Proben liegt der Wert zwischen 0,2 % und 1,0 %. "
        "Diese Zusammenfassung ist ein datenbasierter Vorschlag und fachlich zu pruefen."
    )


def create_review_report(project_id: str, report_docx: Path | None = None) -> str:
    state = load_state(project_id)
    master = state.get("master_samples", [])
    field_unavailable = bool(state.get("field_measurements_unavailable"))
    field_mock = bool(state.get("field_measurements_mock"))
    lines: list[str] = [
        "# Review-Bericht",
        "",
        "Dieser Bericht prueft den aktuellen Projektstand. Warnungen blockieren die weitere Bearbeitung nicht.",
        "",
        "## Weiterarbeit",
        "- Fortfahren moeglich: Ja",
    ]
    if field_unavailable:
        reason = state.get("field_measurements_unavailable", {}).get("reason", "")
        lines.append(f"- Hinweis: {reason}")
        lines.append("- Einschraenkung: Die weitere Bearbeitung ist moeglich, aber Vor-Ort-Felder bleiben offen.")
    elif field_mock:
        reason = state.get("field_measurements_mock", {}).get("reason", "")
        lines.append(f"- WARNUNG: {reason}")
        lines.append("- Einschraenkung: Mockdaten muessen vor fachlicher Verwendung ersetzt werden.")
    else:
        lines.append("- Hinweis: Vor-Ort-Daten wurden nicht als fehlend markiert.")

    lines.extend(
        [
            "",
            "## Probenpruefung",
        ]
    )
    if not master:
        lines.append("- FEHLER: Keine Master-Probenliste vorhanden.")
    for sample in master:
        if sample.get("status") != "ok":
            lines.append(f"- WARNUNG: {sample.get('probe')}: {sample.get('status')}")

    lines.extend(["", "## Tabellenpruefung"])
    counts = summary_counts(master)
    lines.append(f"- Proben gesamt: {counts['total']}")
    lines.append(f"- Chlorid > 1,0 %: {counts['chloride_over_1']}")
    lines.append(f"- Chlorid 0,2 % bis 1,0 %: {counts['chloride_between_02_1']}")
    lines.append(f"- Pruefpunkte offen: {counts['needs_review']}")

    lines.extend(["", "## Word-Pruefung"])
    if report_docx and report_docx.exists():
        text = "\n".join(p.text for p in Document(str(report_docx)).paragraphs)
        if "{{" in text and "}}" in text:
            lines.append("- FEHLER: Im Word-Dokument sind noch Platzhalter vorhanden.")
        else:
            lines.append("- OK: Keine offensichtlichen Textplatzhalter in Absaetzen gefunden.")
    else:
        lines.append("- WARNUNG: Noch kein Word-Entwurf vorhanden.")

    lines.extend(
        [
            "",
            "## SV-Freigabe",
            "- Offen: Bewertungsvorschlaege und fachliche Schlussfolgerungen muessen durch den SV bestaetigt werden.",
            "- Blockiert die weitere Bearbeitung im Prototyp nicht.",
        ]
    )
    return "\n".join(lines) + "\n"


def generate_review(project_id: str) -> dict[str, Any]:
    state = load_state(project_id)
    report_docx = None
    if state.get("outputs", {}).get("report_docx"):
        candidate = project_dir(project_id) / state["outputs"]["report_docx"]
        if candidate.exists():
            report_docx = candidate
    review_path = project_dir(project_id) / "06_review" / "review_report.md"
    review_path.parent.mkdir(parents=True, exist_ok=True)
    review_path.write_text(create_review_report(project_id, report_docx), encoding="utf-8")
    state.setdefault("outputs", {})["review_report"] = str(review_path.relative_to(project_dir(project_id)).as_posix())
    state["status"] = "Review-Bericht erzeugt"
    append_log(state, "Review-Bericht erzeugt; Warnungen blockieren die weitere Bearbeitung nicht.")
    save_state(project_id, state)
    return state


def generate_report(project_id: str, template_relative_path: str | None = None) -> dict[str, Any]:
    state = load_state(project_id)
    output = project_dir(project_id) / "05_report" / "Gutachten_Entwurf.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    result: dict[str, Any] = {}
    if template_relative_path:
        template = resolve_project_file(project_id, template_relative_path)
        result = create_from_template(template, output, state)
        state["selected_files"]["template_file"] = template_relative_path
    else:
        create_from_scratch(output, state)
        result = {"created_without_template": True}

    review_path = project_dir(project_id) / "06_review" / "review_report.md"
    review_path.parent.mkdir(parents=True, exist_ok=True)
    review_path.write_text(create_review_report(project_id, output), encoding="utf-8")

    state["outputs"]["report_docx"] = str(output.relative_to(project_dir(project_id)).as_posix())
    state["outputs"]["review_report"] = str(review_path.relative_to(project_dir(project_id)).as_posix())
    state["report_generation"] = result
    state["status"] = "Gutachtenentwurf erzeugt"
    append_log(state, "Gutachtenentwurf und Review-Bericht erzeugt.")
    save_state(project_id, state)
    return state
